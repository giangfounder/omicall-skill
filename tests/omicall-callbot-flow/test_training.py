"""gen_training.py: ngưỡng lõi câu 30, cache file mẫu khi mạng bị chặn, loại đoạn thực thể lặp sớm."""
import re
import unittest

from _util import SPEC_TRAIN, dump, run, tmpdir

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def make_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "Câu khách nói", "Tên ý định", "Mã ý định"
    ws["A2"] = "(mẫu giả cho test)"
    wb.save(path)
    return str(path)


@unittest.skipIf(openpyxl is None, "SKIP: thiếu openpyxl")
class TestTraining(unittest.TestCase):
    def setUp(self):
        self.d = tmpdir(self)
        self.tpl = make_template(self.d / "tpl.xlsx")

    def test_warn_below_30_cores(self):
        code, so, se = run("gen_training.py", SPEC_TRAIN, "--out", self.d / "o.xlsx", "--template", self.tpl, home=self.d)
        self.assertEqual(code, 0, se)
        self.assertIn("(<30)", so)
        self.assertNotIn("(<25)", so)

    def test_cache_used_when_network_blocked(self):
        home = self.d / "home"
        cache = home / ".cache" / "omicall-skill"
        cache.mkdir(parents=True)
        make_template(cache / "Import_training_data_template.xlsx")
        code, so, se = run("gen_training.py", SPEC_TRAIN, "--out", self.d / "o.xlsx", home=home)
        self.assertEqual(code, 0, se)
        self.assertIn("file mẫu: cache", so)

    def test_no_cache_no_network_clear_error(self):
        home = self.d / "empty_home"
        home.mkdir()
        code, so, se = run("gen_training.py", SPEC_TRAIN, "--out", self.d / "o.xlsx", home=home)
        self.assertNotEqual(code, 0)
        self.assertIn("--template", se)
        self.assertNotIn("Traceback", se)

    def test_early_entity_span_dropped(self):
        spec = {
            "seed": 3, "min_per_intent": 10, "min_per_entity": 40, "noisy_ratio": 0.0, "augment_existing_en": False,
            "intents": {"omi_intent_promise": {"name": "Câu đưa ra lời hẹn", "cores": []}},
            "entities": {"omi_entity_datetime": {"name": "Thời gian", "intent": "omi_intent_promise", "kind": "s",
                                                 "values": ["mai", "thứ hai", "chiều nay", "tuần sau"],
                                                 "patterns": ["{v} hả để {v} hay [{v}] {s} qua", "để [{v}] {s} ghé",
                                                              "hẹn [{v}] nha", "{s} qua vào [{v}]"]}},
        }
        out = self.d / "e.xlsx"
        code, so, se = run("gen_training.py", dump(spec, self.d / "spec.json"), "--out", out, "--template", self.tpl,
                           home=self.d)
        self.assertEqual(code, 0, se)
        m = re.search(r"loại do đoạn thực thể lặp sớm hơn vị trí đánh dấu: (\d+)", so)
        self.assertTrue(m and int(m.group(1)) > 0, so)
        ws = openpyxl.load_workbook(out).worksheets[0]
        rows = [r for r in ws.iter_rows(min_row=3, max_col=6, values_only=True) if r[0]]
        self.assertTrue(rows)
        for r in rows:
            self.assertEqual(r[0].count(r[5]), 1, r)  # đoạn thực thể chỉ xuất hiện một lần → start_char đúng


if __name__ == "__main__":
    unittest.main()
