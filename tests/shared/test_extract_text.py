"""Test extract_text.py: tạo docx (1 bảng), xlsx (2 sheet), PDF 2 trang (1 trang trắng) trong thư mục tạm.

Thiếu python-docx / openpyxl / pymupdf thì test tương ứng SKIP. Không gọi mạng.
"""
import importlib.util
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "shared" / "scripts" / "extract_text.py"


def has(mod):
    return importlib.util.find_spec(mod) is not None


HAS_PDF = has("pymupdf") or has("fitz")


class TestExtractText(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="omi_extract_"))
        self.out = self.tmp / "out" / "source.txt"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def run_x(self, *files):
        p = subprocess.run([sys.executable, "-B", str(SCRIPT), *map(str, files), "--out", str(self.out)],
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        text = self.out.read_text(encoding="utf-8") if self.out.exists() else ""
        return p.returncode, p.stdout + p.stderr, text

    def make_docx(self):
        import docx
        d = docx.Document()
        d.add_paragraph("Đoạn mở đầu DOCX-ALPHA")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Cột A"
        t.cell(0, 1).text = "Cột B"
        t.cell(1, 0).text = "Ô-BETA"
        t.cell(1, 1).text = "Ô-GAMMA"
        d.add_paragraph("Đoạn kết DOCX-OMEGA")
        p = self.tmp / "nghiep_vu.docx"
        d.save(str(p))
        return p

    def make_xlsx(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "KichBan"
        ws.append(["Câu hỏi", "Trả lời"])
        ws.append(["Hỏi XLSX-1", "Đáp XLSX-1"])
        ws2 = wb.create_sheet("BangGia")
        ws2.append(["Gói", "Giá"])
        ws2.append(["Gói XLSX-2", 100])
        p = self.tmp / "du_lieu.xlsx"
        wb.save(str(p))
        return p

    def make_pdf(self):
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "PDF-PAGE-ONE noi dung")
        doc.new_page()  # trang trắng
        p = self.tmp / "tai_lieu.pdf"
        doc.save(str(p))
        doc.close()
        return p

    @unittest.skipUnless(has("docx"), "thiếu python-docx (pip install python-docx)")
    def test_docx_table(self):
        code, out, text = self.run_x(self.make_docx())
        self.assertEqual(code, 0, out)
        for s in ("DOCX-ALPHA", "Ô-BETA", "Ô-GAMMA", "DOCX-OMEGA", "--- bảng 1 ---"):
            self.assertIn(s, text)
        self.assertIn("1 bảng", out)
        self.assertLess(text.index("DOCX-ALPHA"), text.index("Ô-BETA"))
        self.assertLess(text.index("Ô-GAMMA"), text.index("DOCX-OMEGA"))

    @unittest.skipUnless(has("openpyxl"), "thiếu openpyxl (pip install openpyxl)")
    def test_xlsx_two_sheets(self):
        code, out, text = self.run_x(self.make_xlsx())
        self.assertEqual(code, 0, out)
        for s in ("--- sheet KichBan ---", "--- sheet BangGia ---", "Câu hỏi | Trả lời", "Hỏi XLSX-1", "Gói XLSX-2 | 100"):
            self.assertIn(s, text)
        self.assertIn("2 sheet", out)

    @unittest.skipUnless(HAS_PDF, "thiếu pymupdf (pip install pymupdf)")
    def test_pdf_blank_page_warn(self):
        code, out, text = self.run_x(self.make_pdf())
        self.assertEqual(code, 0, out)
        self.assertIn("PDF-PAGE-ONE", text)
        self.assertIn("2 trang", out)
        self.assertIn("WARN tai_lieu.pdf trang 2", out)
        self.assertIn("cần OCR", out)
        self.assertNotIn("WARN tai_lieu.pdf trang 1", out)

    @unittest.skipUnless(has("docx") and has("openpyxl") and HAS_PDF, "thiếu thư viện đọc tài liệu")
    def test_many_files_one_output(self):
        txt = self.tmp / "ghi_chu.md"
        txt.write_text("# Ghi chú MD-DELTA\n", encoding="utf-8")
        code, out, text = self.run_x(self.make_docx(), self.make_xlsx(), self.make_pdf(), txt)
        self.assertEqual(code, 0, out)
        for s in ("DOCX-ALPHA", "XLSX-1", "PDF-PAGE-ONE", "MD-DELTA"):
            self.assertIn(s, text)
        self.assertEqual(out.count("\nOK ") + out.startswith("OK "), 4, out)

    def test_unknown_and_missing(self):
        bad = self.tmp / "x.doc"
        bad.write_bytes(b"\xd0\xcf\x11\xe0")
        code, out, _ = self.run_x(bad, self.tmp / "khong_co.pdf")
        self.assertEqual(code, 1, out)
        self.assertIn("chưa hỗ trợ", out)
        self.assertIn("không thấy file", out)

    def test_all_inputs_fail_no_output(self):
        bad = self.tmp / "x.xyz"
        bad.write_text("abc", encoding="utf-8")
        code, out, text = self.run_x(bad, self.tmp / "khong_co.pdf")
        self.assertEqual(code, 1, out)
        self.assertFalse(self.out.exists(), out)
        self.assertIn("Không đọc được file nào", out)
        self.assertNotIn("Đã ghi", out)

    def test_out_required(self):
        p = subprocess.run([sys.executable, "-B", str(SCRIPT), "a.txt"], capture_output=True, text=True,
                           encoding="utf-8", errors="replace")
        self.assertEqual(p.returncode, 2)


if __name__ == "__main__":
    unittest.main()
