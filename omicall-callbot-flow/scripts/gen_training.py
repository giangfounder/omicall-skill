"""Sinh dữ liệu huấn luyện Kho tri thức (Callbot NORMAL) từ file spec JSON — dùng chung mọi dự án.

  python gen_training.py spec_training.json --out du_lieu.xlsx [--existing da_day_1.xlsx ...] [--template mau.xlsx]

- Mỗi ý định (phần câu không thực thể) và mỗi thực thể: >= min_per_intent / min_per_entity câu MỚI (mặc định 200).
- Câu giống chữ STT tổng đài (stt_noise.py): phiên âm tiếng Anh/tên riêng, nghe nhầm, từ đệm, cắt câu.
- Lọc trùng với các file đã đẩy (--existing) và trong chính file.
- Loại câu có đoạn thực thể xuất hiện sớm hơn vị trí đánh dấu (hệ thống lấy lần xuất hiện đầu → start_char sai), báo số câu bị loại.
- Xuất đúng mẫu Excel hệ thống: --template ghi đè; không có thì dùng bản cache
  ~/.cache/omicall-skill/Import_training_data_template.xlsx (tải từ CDN khi chưa có hoặc để làm mới; mạng hỏng mà có
  cache thì dùng cache; env NO_NETWORK=1 thì không tải). In báo cáo kiểm tra + câu mẫu ngẫu nhiên.
Cấu trúc spec: xem references/spec-training.example.json và references/data-rules.md.
"""
import argparse
import json
import os
import random
import re
import shutil
import sys
import tempfile
import urllib.request
from collections import Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Thiếu openpyxl — cài: python -m pip install openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent))
import stt_noise  # noqa: E402

TEMPLATE_URL = "https://cdn.omicrm.com/callbot/Import_training_data_template.xlsx"
MIN_CORES, MIN_PATTERNS = 30, 15  # khớp SKILL.md Bước 3 và data-rules.md §4
DEFAULT_PERSONAS = [["anh", "em"], ["chị", "em"], ["tui", "em"], ["tôi", "em"], ["chú", "con"], ["cô", "con"]]
OPEN = {
    "q": ["", "", "", "{b} ơi", "cho {s} hỏi", "ủa", "à", "mà", "vậy", "khoan", "{b} cho {s} hỏi cái", "ờ mà"],
    "s": ["", "", "", "{b} ơi", "dạ", "à", "ờ", "ừ", "ừm"],
}
TAIL = {
    "q": ["", "", "", "vậy {b}", "{b}", "hả {b}", "ha {b}", "vậy", "{b} ơi", "đó {b}"],
    "s": ["", "", "", "nha {b}", "{b}", "nghen {b}", "đó {b}", "nha", "nghe {b}"],
}
PARTICLES = {"nha", "nghen", "đi", "vậy", "hả", "ha", "nè", "sau", "rồi", "không", "chưa", "đó", "nghe", "luôn"}
SOFT = {"nha", "nghen", "nghe"}
LEAD = {"ừ", "dạ", "ờ", "ừm", "à", "ủa", "thôi", "hả", "alo", "không", "hông", "chưa"}
SELF = {"anh", "chị", "cô", "chú", "tui", "tôi", "em"}
TEXTING = re.compile(r"\b(ko|k|dc|đc|vs|j|z|r|bn|nt|sdt|sđt|đt)\b")


def clean(t):
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s\[\]]", " ", t.lower())).strip()


def fill_slots(text, slots):
    for name, vals in slots.items():
        while "{" + name + "}" in text:
            text = text.replace("{" + name + "}", random.choice(vals), 1)
    return text


def dress(core, s, b, kind, opens=None, allow_open=True):
    core = core.replace("{s}", s).replace("{b}", b)
    words = core.split()
    op = random.choice(opens or OPEN[kind]).replace("{s}", s).replace("{b}", b) if allow_open and words[0] not in LEAD else ""
    if b in op.split() and b in words:
        op = ""
    tail = ""
    for _ in range(30):
        cand = random.choice(TAIL[kind]).replace("{s}", s).replace("{b}", b)
        last = words[-1].strip("[]")
        if (b in op.split() or b in [w.strip("[]") for w in words]) and b in cand.split():
            continue
        if cand and cand.split()[0] == last:
            continue
        if SOFT & set(w.strip("[]") for w in words) and SOFT & set(cand.split()):
            continue
        if last in PARTICLES and cand and cand.split()[0] != b:
            continue
        if last in SELF and cand == b:
            continue
        tail = cand
        break
    return " ".join(x for x in [op, core, tail] if x)


def split_span(text):
    m = re.search(r"\[(.+?)\]", text)
    if not m:
        return text, None
    return text.replace("[", "").replace("]", ""), m.group(1)


def span_repeats_early(marked):
    """True khi đoạn thực thể [..] đã xuất hiện trước vị trí đánh dấu (sau khi chuẩn hoá như clean())."""
    m = re.search(r"\[(.+?)\]", marked)
    if not m:
        return False
    span = clean(m.group(1))
    if not span:
        return False
    prefix = clean(marked[:m.start()].replace("[", "").replace("]", ""))
    want = len(prefix) + 1 if prefix else 0
    pos = clean(marked.replace("[", "").replace("]", "")).find(span)
    return 0 <= pos < want


def template_path(explicit):
    """--template > cache ~/.cache/omicall-skill/ (làm mới từ CDN khi được) > lỗi rõ."""
    if explicit:
        if not Path(explicit).is_file():
            sys.exit(f"Không thấy file mẫu --template {explicit}")
        return explicit, "--template"
    cache = Path.home() / ".cache" / "omicall-skill" / "Import_training_data_template.xlsx"
    if os.environ.get("NO_NETWORK") == "1":
        if cache.is_file():
            return str(cache), "cache (NO_NETWORK=1)"
        sys.exit(f"NO_NETWORK=1 và chưa có cache {cache} — truyền --template <file mẫu>")
    try:
        cache.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(cache.parent))
        os.close(fd)
        with urllib.request.urlopen(TEMPLATE_URL, timeout=20) as r, open(tmp, "wb") as f:
            shutil.copyfileobj(r, f)
        os.replace(tmp, cache)
        return str(cache), "CDN (đã lưu cache)"
    except Exception as e:  # mạng hỏng / bị chặn
        try:
            os.remove(tmp)
        except Exception:
            pass
        if cache.is_file():
            return str(cache), f"cache (không tải được: {type(e).__name__})"
        sys.exit(f"Không tải được file mẫu ({e}) và chưa có cache {cache} — truyền --template <file mẫu>")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("--out", required=True)
    ap.add_argument("--existing", nargs="*", default=[])
    ap.add_argument("--template")
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding="utf-8"))
    random.seed(spec.get("seed", 1))
    stt_noise.extend(spec.get("en_vi_extra"), spec.get("confuse_extra"), spec.get("fillers_extra"))
    personas = spec.get("personas", DEFAULT_PERSONAS)
    slots = spec.get("slots", {})
    p_noisy = spec.get("noisy_ratio", 0.65)
    n_int, n_ent = spec.get("min_per_intent", 200), spec.get("min_per_entity", 200)
    intents = spec["intents"]
    entities = spec.get("entities", {})
    warnings = []

    template, template_src = template_path(a.template)

    seen, existing_rows = set(), []
    for f in a.existing:
        ws = load_workbook(f).worksheets[0]
        for row in ws.iter_rows(min_row=3, max_col=7, values_only=True):
            if row[0]:
                seen.add(clean(row[0]))
                existing_rows.append(row)
    n_existing = len(seen)
    rows = []
    dropped = Counter()

    def add(text, icode, ecode=None, span=None, marked=None):
        if marked is not None and span_repeats_early(marked):
            dropped[ecode] += 1
            return 0
        text = clean(text)
        if not text or text in seen or TEXTING.search(text):
            return 0
        if ecode:
            span = clean(span or "")
            if not span or span not in text:
                return 0
        seen.add(text)
        rows.append((text, icode, ecode, span))
        return 1

    def fill(n, make, label):
        got = tries = 0
        while got < n and tries < 200000:
            tries += 1
            got += make(tries)
        if got < n:
            warnings.append(f"{label}: chỉ sinh được {got}/{n} câu không trùng — thêm lõi câu/mẫu câu/slot")
        return got

    # ý định có lõi câu riêng (không thực thể)
    for icode, it in intents.items():
        cores = it.get("cores") or []
        if not cores:
            continue
        if len(cores) < MIN_CORES:
            warnings.append(f"{icode}: {len(cores)} lõi câu (<{MIN_CORES}) — câu sẽ ít đa dạng")
        kind = it.get("kind", "s")

        def mk(i, cores=cores, icode=icode, kind=kind, it=it):
            s, b = random.choice(personas)
            core = fill_slots(cores[i % len(cores)], slots)
            allow = random.random() < it.get("open_ratio", 1.0)
            return add(stt_noise.stt(dress(core, s, b, kind, it.get("openers"), allow), random.random() < p_noisy), icode)
        fill(n_int, mk, icode)

    # thực thể: mẫu câu có [đoạn], {x} chủ ngữ, {v} giá trị
    for ecode, en in entities.items():
        pats = en["patterns"]
        if len(pats) < MIN_PATTERNS:
            warnings.append(f"{ecode}: {len(pats)} mẫu câu (<{MIN_PATTERNS}) — câu sẽ ít đa dạng")
        icode, kind = en["intent"], en.get("kind", "q")
        if icode not in intents:
            warnings.append(f"{ecode}: ý định {icode} chưa khai báo trong intents")

        def mk(i, en=en, ecode=ecode, icode=icode, kind=kind, pats=pats):
            s, b = random.choice(personas)
            pat = pats[i % len(pats)]
            if "{x}" in pat:
                pat = pat.replace("{x}", random.choice(en["x"]))
            if "{v}" in pat:
                pat = pat.replace("{v}", en["values"][i % len(en["values"])])
            marked = dress(fill_slots(pat, slots), s, b, kind, en.get("openers"), random.random() < en.get("open_ratio", 1.0))
            noisy = stt_noise.stt(marked, random.random() < p_noisy)
            text, span = split_span(noisy)
            return add(text, icode, ecode, span, noisy)
        fill(n_ent, mk, ecode)

    # biến thể phiên âm cho câu đã đẩy còn từ tiếng Anh viết chuẩn
    n_aug = 0
    if spec.get("augment_existing_en", True):
        for row in existing_rows:
            text1, _, icode1, _, ecode1, span1 = row[:6]
            if not text1 or not stt_noise.has_english(text1) or icode1 not in intents:
                continue
            for _ in range(2):
                if ecode1 and span1 and span1 in text1:
                    i0 = text1.index(span1)
                    noisy = stt_noise.stt(text1[:i0] + "[" + span1 + "]" + text1[i0 + len(span1):], random.random() < p_noisy)
                    t2, sp2 = split_span(noisy)
                    n_aug += add(t2, icode1, ecode1, sp2, noisy)
                else:
                    n_aug += add(stt_noise.stt(text1, random.random() < p_noisy), icode1)

    wb = load_workbook(template)
    ws = wb.worksheets[0]
    if ws.max_row > 2:
        ws.delete_rows(3, ws.max_row)
    for r, (text, icode, ecode, span) in enumerate(rows, start=3):
        vals = [text, intents[icode]["name"], icode]
        if ecode:
            en = entities.get(ecode, {})
            vals += [en.get("name", ecode), ecode, span, en.get("type", "STRING")]
        for c, v in enumerate(vals, start=1):
            ws.cell(r, c, v).alignment = Alignment(vertical="center")
    for idx, col in enumerate(ws.iter_cols(min_row=2), start=1):
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(w + 2, 10), 80)
    wb.save(a.out)

    texts = [r[0] for r in rows]
    english_left = sum(stt_noise.has_english(t) for t in texts)
    print(f"{len(rows)} câu mới -> {a.out} (lọc trùng với {n_existing} câu đã có; biến thể phiên âm từ câu cũ: {n_aug})")
    print(f"file mẫu: {template_src}")
    print(f"loại do đoạn thực thể lặp sớm hơn vị trí đánh dấu: {sum(dropped.values())}", dict(dropped) if dropped else "")
    print("ý định:", dict(Counter(r[1] for r in rows)))
    print("thực thể:", dict(Counter(r[2] for r in rows if r[2])))
    print("kiểm tra: trùng", len(texts) - len(set(texts)), "| dấu câu", sum(bool(re.search(r"[^\w\s]", t)) for t in texts),
          "| đoạn thực thể sai", sum(1 for r in rows if r[2] and r[3] not in r[0]), "| còn tiếng Anh", english_left)
    for w in warnings:
        print("CẢNH BÁO:", w)
    rnd = random.Random(7)
    for key in list(intents) + list(entities):
        pool = [r for r in rows if (r[1] == key and not r[2]) or r[2] == key]
        if pool:
            print(f"\n# {key}")
            for r in rnd.sample(pool, min(5, len(pool))):
                print("   ", r[0], ("| " + r[3]) if r[3] else "")


if __name__ == "__main__":
    main()
